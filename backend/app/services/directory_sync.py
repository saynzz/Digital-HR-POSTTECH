import csv
import io
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta, timezone

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models import AccountStatus, Employee, ImportRun, Role, Track, User
from app.services.notifications import notify

FIELDS = {
    "external_id": ["табельныйномер", "табельный", "табномер", "таб№", "№", "код", "кодсотрудника", "externalid", "id",
                    "employeeid", "tabnumber"],
    "full_name": ["фио", "сотрудник", "фамилияимяотчество", "работник", "fullname", "name"],
    "email": ["email", "email", "почта", "электроннаяпочта", "рабочаяпочта", "mail"],
    "phone": ["телефон", "мобильныйтелефон", "phone"],
    "position": ["должность", "position", "title"],
    "department": ["подразделение", "отдел", "департамент", "department"],
    "manager_external_id": ["табельныйномерруководителя", "таб№руководителя", "кодруководителя", "руководительтаб№",
                            "руководителькод", "managerid", "managerexternalid"],
    "manager_name": ["руководитель", "руководительфио", "manager", "managername"],
    "hire_date": ["датаприема", "датаприёма", "датаприеманаработу", "дататрудоустройства", "hiredate", "startdate"],
    "vacation_balance_days": ["остатокотпуска", "остатокотпускадн", "остатокотпускадней", "отпуск", "vacationbalance",
                              "vacationdays", "vacationbalancedays"],
    "role": ["роль", "role"],
    "dismissal_date": ["датаувольнения", "уволен", "dismissaldate", "terminationdate"],
}
REQUIRED = ["external_id", "full_name", "position", "department", "hire_date"]
LABELS = {"external_id": "Табельный номер", "full_name": "ФИО", "position": "Должность", "department": "Подразделение",
          "hire_date": "Дата приёма", "email": "Email", "phone": "Телефон", "manager_external_id": "Руководитель",
          "vacation_balance_days": "Остаток отпуска", "role": "Роль"}
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _norm(h: str) -> str:
    return re.sub(r"[\s\.\,\-_:;()\"'/]", "", str(h or "").strip().lower().replace("ё", "е")).replace("e-mail", "email")


ALIASES = {_norm(a): k for k, al in FIELDS.items() for a in al}


@dataclass
class Record:
    row: int
    external_id: str = ""
    full_name: str = ""
    email: str | None = None
    phone: str | None = None
    position: str = ""
    department: str = ""
    manager_external_id: str | None = None
    manager_name: str | None = None
    hire_date: str | None = None
    vacation_balance_days: int | None = None
    role: str | None = None
    dismissal_date: str | None = None
    errors: list[str] = field(default_factory=list)


def _date(v) -> str | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y", "%Y-%m-%dT%H:%M:%S", "%d.%m.%Y %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(s)


def _clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _role(v: str) -> str | None:
    s = v.lower()
    if not s:
        return None
    if s.startswith("рук") or s.startswith("man") or s in ("начальник", "директор"):
        return "manager"
    if s.startswith("сотр") or s.startswith("emp") or s.startswith("раб"):
        return "employee"
    return "invalid"


def to_records(rows: list[dict], start_row: int = 2) -> list[Record]:
    out = []
    for i, raw in enumerate(rows):
        r = Record(row=start_row + i)
        mapped = {}
        for k, v in raw.items():
            key = ALIASES.get(_norm(k))
            if key and key not in mapped:
                mapped[key] = v
        if not any(_clean(v) for v in mapped.values()):
            continue
        r.external_id = _clean(mapped.get("external_id"))
        r.full_name = _clean(mapped.get("full_name"))
        r.position = _clean(mapped.get("position"))
        r.department = _clean(mapped.get("department"))
        r.email = _clean(mapped.get("email")).lower() or None
        r.phone = _clean(mapped.get("phone")) or None
        r.manager_external_id = _clean(mapped.get("manager_external_id")) or None
        r.manager_name = _clean(mapped.get("manager_name")) or None
        for key in ("hire_date", "dismissal_date"):
            try:
                setattr(r, key, _date(mapped.get(key)))
            except ValueError as e:
                r.errors.append(f"Не удалось прочитать дату «{e}»")
        bal = _clean(mapped.get("vacation_balance_days"))
        if bal:
            try:
                r.vacation_balance_days = int(float(bal.replace(",", ".")))
            except ValueError:
                r.errors.append(f"Остаток отпуска должен быть числом: «{bal}»")
        role = _role(_clean(mapped.get("role")))
        if role == "invalid":
            r.errors.append("Роль: укажите «сотрудник» или «руководитель»")
        else:
            r.role = role
        for f in REQUIRED:
            if not getattr(r, f):
                r.errors.append(f"Не заполнено: {LABELS[f]}")
        if r.email and not EMAIL_RE.match(r.email):
            r.errors.append(f"Некорректный email: {r.email}")
        out.append(r)
    return out


def read_file(filename: str, content: bytes) -> list[Record]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = None
        start = 1
        for row in it:
            start += 1
            if row and sum(1 for c in row if ALIASES.get(_norm(c or ""))) >= 2:
                header = [str(c or "") for c in row]
                break
        if header is None:
            raise ValueError("Не нашли строку заголовков: нужны колонки «Табельный номер», «ФИО», «Должность» и т.д.")
        rows = [dict(zip(header, r)) for r in it]
        return to_records(rows, start)
    if name.endswith((".csv", ".txt")):
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("Не удалось определить кодировку файла")
        sample = text[:4096]
        delim = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.DictReader(io.StringIO(text), delimiter=delim)
        if not reader.fieldnames or sum(1 for h in reader.fieldnames if ALIASES.get(_norm(h))) < 2:
            raise ValueError("Не нашли заголовки колонок. Скачайте шаблон и заполните его")
        return to_records(list(reader))
    raise ValueError("Поддерживаются файлы .xlsx и .csv")


def _today() -> date:
    return datetime.now(timezone(timedelta(hours=5))).date()


def build_plan(db: Session, records: list[Record], full_sync: bool = False) -> dict:
    employees = {e.external_id: e for e in db.scalars(select(Employee))}
    users_by_emp = {u.employee_id: u for u in db.scalars(select(User))}
    users_by_email = {u.email.lower(): u for u in users_by_emp.values()}
    seen_ids: dict[str, int] = {}
    seen_emails: dict[str, int] = {}
    file_ids = {r.external_id for r in records if r.external_id}
    names = {r.full_name.lower(): r.external_id for r in records if r.full_name}
    names.update({e.full_name.lower(): e.external_id for e in employees.values()})
    manager_ids = set()
    for r in records:
        if not r.manager_external_id and r.manager_name:
            key = r.manager_name.lower()
            match = names.get(key) or next((v for k, v in names.items() if k.startswith(key)), None)
            if match:
                r.manager_external_id = match
            else:
                r.errors.append(f"Руководитель «{r.manager_name}» не найден")
        if r.manager_external_id:
            manager_ids.add(r.manager_external_id)

    items = []
    today = _today().isoformat()
    for r in records:
        if r.external_id in seen_ids:
            r.errors.append(f"Табельный номер повторяется (строка {seen_ids[r.external_id]})")
        seen_ids.setdefault(r.external_id, r.row)
        if r.email:
            if r.email in seen_emails:
                r.errors.append(f"Email повторяется (строка {seen_emails[r.email]})")
            seen_emails.setdefault(r.email, r.row)
            owner = users_by_email.get(r.email)
            emp = employees.get(r.external_id)
            if owner and (emp is None or owner.employee_id != emp.id):
                r.errors.append(f"Email {r.email} уже занят другим сотрудником")
        if r.manager_external_id:
            if r.manager_external_id == r.external_id:
                r.errors.append("Сотрудник указан руководителем самому себе")
            elif r.manager_external_id not in file_ids and r.manager_external_id not in employees:
                r.errors.append(f"Руководитель с табельным номером {r.manager_external_id} не найден")
        role = r.role or ("manager" if r.external_id in manager_ids else "employee")
        dismissed = bool(r.dismissal_date and r.dismissal_date <= today)
        emp = employees.get(r.external_id)
        user = users_by_emp.get(emp.id) if emp else None
        item = {**{k: v for k, v in asdict(r).items() if k != "errors"}, "role": role, "errors": r.errors,
                "changes": [], "action": "error" if r.errors else "skip", "account": None}
        if not r.errors:
            if dismissed:
                item["action"] = "dismiss" if user and user.status != AccountStatus.blocked else "skip"
            elif emp is None:
                item["action"] = "create"
            else:
                diff = []
                for f, cur in (("full_name", emp.full_name), ("position", emp.position), ("department", emp.department),
                               ("hire_date", emp.hire_date.isoformat())):
                    if getattr(r, f) != cur:
                        diff.append({"field": LABELS[f], "from": cur, "to": getattr(r, f)})
                cur_mgr = next((x.external_id for x in employees.values() if x.id == emp.manager_id), None)
                if (r.manager_external_id or None) != cur_mgr:
                    diff.append({"field": "Руководитель", "from": cur_mgr or "—", "to": r.manager_external_id or "—"})
                if r.vacation_balance_days is not None and r.vacation_balance_days != emp.vacation_balance_days:
                    diff.append({"field": "Остаток отпуска", "from": emp.vacation_balance_days, "to": r.vacation_balance_days})
                if r.email and user and user.email.lower() != r.email:
                    diff.append({"field": "Email", "from": user.email, "to": r.email})
                item["changes"] = diff
                item["action"] = "update" if diff else "skip"
            if not dismissed:
                if user is None:
                    item["account"] = "invite" if r.email else "no_email"
                elif user.role == Role.hr:
                    item["role"] = "hr"
                    item["account"] = "exists"
                else:
                    item["account"] = "exists"
        items.append(item)

    missing = []
    if full_sync:
        for ext, e in employees.items():
            u = users_by_emp.get(e.id)
            if ext not in file_ids and not (u and u.role == Role.hr) and not (u and u.status == AccountStatus.blocked):
                missing.append({"external_id": ext, "full_name": e.full_name, "department": e.department})
    count = lambda a: sum(1 for i in items if i["action"] == a)
    summary = {"total": len(items), "create": count("create"), "update": count("update"), "dismiss": count("dismiss"),
               "skip": count("skip"), "error": count("error"), "missing": len(missing),
               "invites": sum(1 for i in items if i["account"] == "invite" and i["action"] in ("create", "update", "skip")),
               "no_email": sum(1 for i in items if i["account"] == "no_email")}
    return {"items": items, "missing": missing, "summary": summary}


def create_run(db: Session, user: User | None, source: str, filename: str | None, records: list[Record],
               full_sync: bool = False) -> ImportRun:
    stored = [asdict(r) for r in records]
    plan = build_plan(db, records, full_sync)
    run = ImportRun(source=source, filename=filename, created_by_user_id=user.id if user else None,
                    records=stored, summary={**plan["summary"], "full_sync": full_sync})
    db.add(run)
    db.commit()
    db.refresh(run)
    return run


def _block(db: Session, user: User) -> None:
    from app.models import InviteToken

    user.status = AccountStatus.blocked
    for t in db.scalars(select(InviteToken).where(InviteToken.user_id == user.id, InviteToken.used_at.is_(None),
                                                  InviteToken.revoked_at.is_(None))):
        t.revoked_at = datetime.now(timezone.utc)


def apply_run(db: Session, run: ImportRun, actor: User | None, send_invites: bool, block_missing: bool,
              assign_onboarding: bool) -> dict:
    from app.services import onboarding
    from app.services.gamification import enroll_track

    if run.applied:
        raise ValueError("Этот импорт уже применён")
    records = [Record(**{**r, "errors": list(r.get("errors", []))}) for r in run.records]
    plan = build_plan(db, records, run.summary.get("full_sync", False))
    employees = {e.external_id: e for e in db.scalars(select(Employee))}
    created, updated, dismissed, invited, blocked, mail_errors = [], [], [], [], [], []
    new_ids: set[int] = set()

    for it in plan["items"]:
        if it["action"] in ("error",):
            continue
        e = employees.get(it["external_id"])
        if it["action"] == "create":
            e = Employee(external_id=it["external_id"], full_name=it["full_name"], company=settings.company_name,
                         position=it["position"], department=it["department"],
                         hire_date=date.fromisoformat(it["hire_date"]),
                         vacation_balance_days=it["vacation_balance_days"] or 0)
            db.add(e)
            db.flush()
            employees[e.external_id] = e
            new_ids.add(e.id)
            created.append(e.full_name)
        elif it["action"] == "update" and e is not None:
            e.full_name, e.position, e.department = it["full_name"], it["position"], it["department"]
            e.hire_date = date.fromisoformat(it["hire_date"])
            if it["vacation_balance_days"] is not None:
                e.vacation_balance_days = it["vacation_balance_days"]
            updated.append(e.full_name)

    for it in plan["items"]:
        if it["action"] in ("create", "update", "skip") and not it["errors"]:
            e = employees.get(it["external_id"])
            mgr = employees.get(it["manager_external_id"]) if it["manager_external_id"] else None
            if e is not None:
                e.manager_id = mgr.id if mgr else None
    db.flush()

    users_by_emp = {u.employee_id: u for u in db.scalars(select(User))}
    for it in plan["items"]:
        e = employees.get(it["external_id"])
        if e is None or it["action"] == "error":
            continue
        user = users_by_emp.get(e.id)
        if it["action"] == "dismiss":
            if user and user.status != AccountStatus.blocked:
                _block(db, user)
                dismissed.append(e.full_name)
            continue
        role = Role.manager if it["role"] == "manager" else Role.employee
        fresh = False
        if user is None and it["email"]:
            user = User(email=it["email"], password_hash=None, role=role, status=AccountStatus.invited, employee=e)
            db.add(user)
            db.flush()
            users_by_emp[e.id] = user
            fresh = True
        elif user is not None:
            if it["email"] and user.email.lower() != it["email"]:
                user.email = it["email"]
            if user.role != Role.hr and user.role != role:
                user.role = role
        if user and fresh and send_invites and actor:
            url = onboarding.issue_invite(db, actor, user, commit=False)
            invited.append({"name": e.full_name, "email": user.email, "url": url})
        if assign_onboarding and e.id in new_ids and (_today() - e.hire_date).days <= 60:
            track = db.scalar(select(Track).where(Track.slug == "onboarding"))
            if track:
                enroll_track(db, e, track, actor, _today() + timedelta(days=30))

    if block_missing:
        for m in plan["missing"]:
            e = employees.get(m["external_id"])
            u = users_by_emp.get(e.id) if e else None
            if u and u.status != AccountStatus.blocked and u.role != Role.hr:
                _block(db, u)
                blocked.append(e.full_name)

    db.commit()
    for inv in invited:
        u = db.scalar(select(User).where(func.lower(User.email) == inv["email"].lower()))
        ok, err = onboarding.send_invite_email(u, inv["url"])
        inv["email_sent"] = ok
        if err and err not in mail_errors:
            mail_errors.append(err)
    run.applied = True
    run.applied_at = datetime.now(timezone.utc)
    run.result = {"created": created, "updated": updated, "dismissed": dismissed, "blocked_missing": blocked,
                  "invited": [{"name": i["name"], "email": i["email"], "email_sent": i.get("email_sent", False)} for i in invited],
                  "mail_errors": mail_errors}
    if actor:
        notify(db, actor.id, "system", f"Импорт сотрудников применён: +{len(created)}, изменено {len(updated)}",
               f"Приглашений: {len(invited)}, уволено: {len(dismissed)}", "/hr?tab=import")
    db.commit()
    return run.result


TEMPLATE_HEADER = ["Табельный номер", "ФИО", "Email", "Телефон", "Должность", "Подразделение",
                   "Табельный номер руководителя", "Дата приёма", "Остаток отпуска", "Роль", "Дата увольнения"]
TEMPLATE_ROWS = [
    ["ZUP-1001", "Абдрахманов Саят Нурланович", "sayat.a@company.kz", "+7 701 000 00 11", "Начальник склада",
     "Склад Алматы", "", "01.03.2019", "24", "руководитель", ""],
    ["ZUP-1002", "Иванова Ольга Петровна", "olga.i@company.kz", "", "Кладовщик", "Склад Алматы", "ZUP-1001",
     "15.06.2023", "18", "", ""],
    ["ZUP-1003", "Тулегенов Ержан Маратович", "erzhan.t@company.kz", "", "Водитель-экспедитор", "Склад Алматы",
     "ZUP-1001", "10.09.2024", "12", "", ""],
]


def template_csv() -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(TEMPLATE_HEADER)
    w.writerows(TEMPLATE_ROWS)
    return "﻿" + buf.getvalue()
